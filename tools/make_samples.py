"""Generate synthetic test fixtures; none of the measurements are real research."""
from pathlib import Path
from PIL import Image, ImageDraw
from docx import Document
from pptx import Presentation
from pptx.util import Inches, Pt
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "samples"
OUT.mkdir(exist_ok=True)

reports = {
    "01-ceramic-coating.md": """# 模擬研究 A：陶瓷塗層降低電池溫升
本文件為合成測試資料，並非真實研究結果。
實驗 EXP-A01，研究團隊：Thermal Lab，日期 2025-03-12。
目標：在不降低容量的前提下降低電池充放電溫升。
條件：環境 25°C、放電倍率 2C、相同電芯型號、樣本數 n=6。
控制組未塗層，峰值溫度 48°C，容量保持率 97%。
陶瓷塗層 20 μm，峰值溫度 42°C，容量保持率 96.8%。
結果：在本實驗條件下，陶瓷塗層使峰值溫度降低 6°C。
限制：尚未測試 45°C 環境、長期循環或量產良率，不可推論所有條件都有效。
下一步：測試高溫環境與塗層厚度，確認熱阻和容量衰退。
關聯材料：氧化鋁陶瓷、鋰離子電芯、導熱黏著劑。
""",
    "02-graphite-spreader.md": """# 模擬研究 B：石墨散熱片與熱傳導
本文件為合成測試資料，並非真實研究結果。
實驗 EXP-B07，研究團隊：Packaging Lab，日期 2025-04-08。
條件：環境 35°C、放電倍率 3C、樣本數 n=4。
控制組峰值溫度 61°C，增加 0.3 mm 石墨散熱片後峰值溫度 54°C。
石墨片可改善面內熱擴散；厚度增加會增加重量與封裝體積。
注意：本實驗與 EXP-A01 的環境溫度、倍率及樣本數不同，不能直接認定石墨片優於陶瓷塗層。
後續需要相同電芯、25°C 和 2C 條件的對照試驗。
""",
    "03-adhesive-change.md": """# 模擬變更評估：更換導熱黏著劑
本文件為合成測試資料，並非真實工程結論。
變更 CHG-019：將導熱黏著劑 A 更換為黏著劑 B。
已知關聯：黏著層厚度影響界面熱阻，固化溫度影響電芯封裝製程。
已完成證據：小樣測試中 B 的剝離強度較 A 低 12%，測試條件 25°C、拉伸速率 50 mm/min。
潛在 impact：散熱效率、黏著可靠度、裝配時間與高溫老化；這些項目尚未完成實驗驗證。
相關研究：EXP-A01 陶瓷塗層、EXP-B07 石墨散熱片均使用黏著劑 A。
建議：先評估熱循環後剝離強度，再測界面熱阻。不要把關聯當成已證實因果。
""",
    "04-onboarding.md": """# 模擬 RD 新人實驗指南
此為合成訓練資料。
開始實驗前，確認 experiment ID、文件版本、樣本數、控制組與量測儀器校正日期。
記錄環境溫度、濕度、倍率、材料批次、厚度及單位。
比較兩份研究時必須先對齊實驗條件，不能只比較結果數字。
找不到研究報告僅代表目前知識庫沒有檢索到，不能斷言公司沒有人做過。
主管彙整應交代涵蓋的研究範圍、已證實成果、待驗證風險及下一步實驗。
""",
    "05-english-study.md": """# Synthetic study C: thermal cycling and bonding reliability
SYNTHETIC DATA ONLY. Experiment EXP-C03, Reliability Lab.
The team evaluated adhesive A and B after 100 thermal cycles between -20 and 60 degrees Celsius.
Adhesive B exhibited edge delamination in 2 out of 8 samples; adhesive A showed none in 8 samples.
This observation is not a lifetime prediction. Cell chemistry and production lots were not controlled.
Research opportunity: investigate whether bond-line thickness interacts with graphite heat spreaders.
Compare only after matching sample preparation, curing profile, and measurement conditions.
""",
}
for name, text in reports.items():
    (OUT / name).write_text(text, encoding="utf-8")

image = Image.new("RGB", (1000, 560), "#f4f8fa")
draw = ImageDraw.Draw(image)
draw.text((35, 25), "SYNTHETIC DATA - Peak temperature / EXP-A01", fill="#132b3b", font_size=27)
draw.line((100, 110, 100, 450, 920, 450), fill="#243d4a", width=3)
for i, (name, value) in enumerate((("Control", 48), ("Ceramic coating", 42))):
    x = 240 + i * 350
    y = 450 - value * 6
    draw.rectangle((x, y, x + 170, 450), fill=("#768b9d" if i == 0 else "#218c75"))
    draw.text((x + 40, y - 42), f"{value} C", fill="#182f3d", font_size=28)
    draw.text((x - 5, 470), name, fill="#182f3d", font_size=22)
draw.text((100, 525), "Ambient 25 C | 2C discharge | n=6 | Illustrative, not real results", fill="#536d7d", font_size=18)
chart = OUT / "synthetic-temperature-chart.png"
image.save(chart)
image.save(OUT / "09-synthetic-scanned-chart.pdf", "PDF", resolution=100)

prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[5])
slide.shapes.title.text = "模擬實驗 EXP-A01：陶瓷塗層"
box = slide.shapes.add_textbox(Inches(.5), Inches(1.0), Inches(9), Inches(.7))
box.text_frame.text = "合成測試資料。25°C / 2C / n=6。未塗層 48°C，陶瓷塗層 42°C。"
slide.shapes.add_picture(str(chart), Inches(.8), Inches(1.8), width=Inches(8.3))
slide = prs.slides.add_slide(prs.slide_layouts[1])
slide.shapes.title.text = "限制與下一步"
slide.placeholders[1].text = "尚未驗證高溫環境與長期循環。\n不能直接與 35°C / 3C 的石墨片實驗比較。\n下一步：對齊實驗條件，測試黏著劑 B 的可靠度。"
prs.save(OUT / "06-synthetic-research.pptx")

doc = Document()
doc.add_heading("模擬研究摘要：散熱材料", 0)
doc.add_paragraph("本文件全部為合成資料，供軟體測試。陶瓷塗層在 25°C / 2C 條件下，峰值溫度從 48°C 降為 42°C。")
doc.add_picture(str(chart), width=__import__('docx').shared.Inches(6))
doc.add_paragraph("關聯：更換黏著劑可能影響界面熱阻，必須進行實驗確認。")
doc.save(OUT / "07-synthetic-summary.docx")

wb = Workbook()
ws = wb.active
ws.title = "實驗比較"
for row in [
    ["Synthetic data only", "環境 °C", "倍率", "n", "控制組 °C", "實驗組 °C"],
    ["EXP-A01 陶瓷塗層", 25, "2C", 6, 48, 42],
    ["EXP-B07 石墨散熱片", 35, "3C", 4, 61, 54],
    ["注意：條件不同，不可直接判定優劣"],
]: ws.append(row)
wb.save(OUT / "08-synthetic-results.xlsx")
print(f"Created synthetic fixtures in {OUT}")
