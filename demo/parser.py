"""CPU document adapters retaining source location and embedded pictures.

OCR is optional via the separate CPU enrichment function below.
No slide rendering, formula evaluation, or visual understanding is claimed.
"""
from pathlib import Path
import json
import re
import threading

_ocr = None
_ocr_lock = threading.Lock()


def enrich_ocr(pages, destination):
    global _ocr
    warnings = []
    with _ocr_lock:
        from rapidocr_onnxruntime import RapidOCR
        if _ocr is None:
            _ocr = RapidOCR(intra_op_num_threads=2, inter_op_num_threads=1)
        for page in pages:
            page["ocr"] = []
            for filename in page["images"][:8]:
                try:
                    result, _ = _ocr(str(destination / filename))
                    lines = [{"text": row[1], "confidence": float(row[2])} for row in (result or []) if float(row[2]) >= .5]
                    page["ocr"].append({"image": filename, "lines": lines})
                    if lines:
                        page["text"] += "\n[圖片 OCR 辨識文字，需核對原圖]\n" + "\n".join(x["text"] for x in lines)
                except Exception:
                    warnings.append(f"{page['location']} 的 {filename} OCR 失敗；原圖仍保留。")
            if len(page["images"]) > 8:
                warnings.append(f"{page['location']} 本版僅 OCR 前 8 張圖片。")
    return warnings

SUPPORTED = {".txt", ".md", ".pdf", ".docx", ".pptx", ".xlsx", ".png", ".jpg", ".jpeg"}


def parse_document(path: Path, destination: Path):
    if path.suffix.lower() in {".pptx", ".docx", ".xlsx"}:
        from zipfile import ZipFile
        with ZipFile(path) as archive:
            entries = archive.infolist()
            if len(entries) > 20000 or sum(x.file_size for x in entries) > 256 * 1024 * 1024:
                raise ValueError("Office archive exceeds demo expansion limit")
    pages = []
    counter = 0

    def picture(blob, extension):
        nonlocal counter
        counter += 1
        # Convert to safe browser-renderable PNG, avoiding active SVG/HTML content.
        from PIL import Image
        from io import BytesIO
        try:
            img = Image.open(BytesIO(blob))
            img.thumbnail((1800, 1800))
            filename = f"image-{counter}.png"
            img.convert("RGB").save(destination / filename)
            return filename
        except (OSError, ValueError):
            return None

    def add(label, text, images=None):
        pages.append({"location": label, "text": text.strip(), "images": [x for x in (images or []) if x]})

    suffix = path.suffix.lower()
    if suffix in (".txt", ".md"):
        add("全文", path.read_text(encoding="utf-8-sig"))
    elif suffix == ".pptx":
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        def shapes_text(shapes, texts, images):
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes_text(shape.shapes, texts, images)
                if shape.has_text_frame:
                    texts.append(shape.text)
                if shape.has_table:
                    texts.extend(" | ".join(c.text for c in row.cells) for row in shape.table.rows)
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    images.append(picture(shape.image.blob, shape.image.ext))
                if shape.has_chart:
                    try:
                        texts.append("Chart categories: " + ", ".join(str(c.label) for c in shape.chart.plots[0].categories))
                        for series in shape.chart.series:
                            texts.append(f"{series.name}: {list(series.values)}")
                    except (AttributeError, ValueError, TypeError):
                        texts.append("[Chart values could not be extracted]")
        for number, slide in enumerate(Presentation(path).slides, 1):
            texts, images = [], []
            shapes_text(slide.shapes, texts, images)
            add(f"投影片 {number}", "\n".join(texts), images)
    elif suffix == ".docx":
        from docx import Document
        doc = Document(path)
        texts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            texts.extend(" | ".join(c.text for c in row.cells) for row in table.rows)
        images = []
        for rel in doc.part.rels.values():
            if "image" in rel.reltype and not rel.is_external:
                images.append(picture(rel.target_part.blob, "png"))
        add("全文（未計算 Word 頁碼）", "\n".join(texts), images)
    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=False)
        try:
            for sheet in wb:
                for start in range(1, sheet.max_row + 1, 20):
                    rows = []
                    for row in sheet.iter_rows(min_row=start, max_row=min(start + 19, sheet.max_row)):
                        rows.append(" | ".join(f"{c.coordinate}={c.value}" for c in row if c.value is not None))
                    if any(rows):
                        add(f"{sheet.title} 列 {start}–{min(start+19, sheet.max_row)}", "\n".join(rows))
        finally:
            wb.close()
    elif suffix == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported in this demo.")
        for number, page in enumerate(reader.pages, 1):
            images = []
            for img in page.images:
                images.append(picture(img.data, "png"))
            add(f"PDF 第 {number} 頁", page.extract_text() or "", images)
    elif suffix in (".png", ".jpg", ".jpeg"):
        add("圖片", "", [picture(path.read_bytes(), suffix)])
    else:
        raise ValueError("Unsupported file type")
    if not pages:
        raise ValueError("No readable content")
    return pages
